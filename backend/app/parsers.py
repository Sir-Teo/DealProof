from __future__ import annotations

import csv
import ipaddress
import io
import re
import socket
from pathlib import Path
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from .config import HTTP_USER_AGENT


class UrlFetchError(RuntimeError):
    pass


def validate_fetch_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise UrlFetchError("Failed to fetch URL: URL must start with http:// or https:// and include a hostname.")
    hostname = parsed.hostname.lower()
    if hostname in {"localhost", "localhost.localdomain"}:
        raise UrlFetchError("Failed to fetch URL: localhost and private network URLs are not allowed.")
    try:
        addresses = socket.getaddrinfo(hostname, None, type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise UrlFetchError("Failed to fetch URL: hostname could not be resolved.") from exc
    for address in {item[4][0] for item in addresses}:
        ip = ipaddress.ip_address(address)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_multicast or ip.is_unspecified:
            raise UrlFetchError("Failed to fetch URL: localhost and private network URLs are not allowed.")


def request_headers_for_url(url: str) -> dict[str, str]:
    headers = {
        "User-Agent": HTTP_USER_AGENT,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,text/plain;q=0.8,*/*;q=0.5",
    }
    if "sec.gov" in url.lower():
        headers["Accept-Encoding"] = "gzip, deflate"
    return headers


def summarize(text: str) -> tuple[str, str]:
    clean = re.sub(r"\s+", " ", text).strip()
    if not clean:
        return "No readable text extracted.", ""
    sentences = re.split(r"(?<=[.!?])\s+", clean)
    return " ".join(sentences[:2])[:280], clean[:520]


def _parse_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    image_page_count = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            pages.append(text)
        else:
            image_page_count += 1
    extracted = "\n".join(pages)
    # If more than half the pages yielded no text, the PDF is likely image-based.
    # Surface a clear warning so the user knows to upload a text version or transcript.
    total_pages = len(reader.pages)
    if total_pages > 0 and image_page_count / total_pages > 0.5:
        note = (
            f"[NOTE: {image_page_count} of {total_pages} pages are image-only and could not be read. "
            "For best results, upload the original editable file or a plain-text transcript alongside this PDF.]"
        )
        extracted = f"{note}\n\n{extracted}" if extracted else note
    return extracted


def parse_file(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs)
    if suffix == ".csv":
        data = path.read_text(errors="ignore")
        rows = list(csv.reader(io.StringIO(data)))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    if suffix in {".xlsx", ".xls"}:
        return _parse_xlsx(path)
    if suffix in {".pptx", ".ppt"}:
        return _parse_pptx(path)
    return path.read_text(errors="ignore")


def _parse_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"## {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) if cell is not None else "" for cell in row]
            if any(c.strip() for c in cells):
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _parse_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return ""
    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        slide_parts: list[str] = [f"## Slide {i}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if text:
                    slide_parts.append(text)
            if hasattr(shape, "notes") and shape.notes and shape.notes.text_frame:
                notes = shape.notes.text_frame.text.strip()
                if notes:
                    slide_parts.append(f"[Notes: {notes}]")
        if len(slide_parts) > 1:
            parts.append("\n".join(slide_parts))
    return "\n\n".join(parts)


async def fetch_url_text(url: str) -> str:
    validate_fetch_url(url)
    async with httpx.AsyncClient(timeout=15, follow_redirects=True) as client:
        try:
            response = await client.get(url, headers=request_headers_for_url(url))
            response.raise_for_status()
        except httpx.HTTPStatusError as exc:
            status = exc.response.status_code
            raise UrlFetchError(f"Failed to fetch URL: upstream returned HTTP {status}.") from exc
        except httpx.HTTPError as exc:
            raise UrlFetchError(f"Failed to fetch URL: {exc}") from exc
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else url
    body = soup.get_text(" ", strip=True)
    return f"{title}\n{body}"


def infer_kind(name: str) -> str:
    lower = name.lower()
    if lower.endswith(".pdf") or "deck" in lower:
        return "deck"
    if lower.endswith(".csv") or "financial" in lower:
        return "financials"
    if "transcript" in lower:
        return "transcript"
    return "document"
